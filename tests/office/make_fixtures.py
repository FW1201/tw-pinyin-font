"""產生 Office 實測檔（Word／PowerPoint／Excel），字型設為 Taiwan Pinyin Kai。

    python tests/office/make_fixtures.py [outdir]

Word 檔每個測試項目各有兩段：
  A = Word 預設（不指定 OpenType 功能）
  B = 開啟「使用內容替代字」（w14:cntxtAlts），看 rclt／calt 是否需要手動開啟
另產生 embed.docx：設定「在檔案中內嵌字型」，用來測內嵌。
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font as XlFont
from pptx import Presentation
from pptx.util import Inches, Pt as PPt

FONT = "Taiwan Pinyin Kai"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
CASES = [
    ("預設讀音", "長"),
    ("詞語規則", "長大　銀行　目的　睡著"),
    ("IVS 指定", "銀行\U000E01E0　長\U000E01E1　行\U000E01E2"),
    ("段落", "我們去銀行存錢，弟弟長大以後想當醫生。音樂課的時候，大家都很快樂。"),
]
EXPECT = {
    "預設讀音": "cháng",
    "詞語規則": "zhǎng dà／yín háng／mù dì／shuì zháo",
    "IVS 指定": "yín xíng／zhǎng／háng",
    "段落": "yín háng、dì di、zhǎng、yuè、lè",
}


def set_run_font(run, size=28, contextual=False):
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    if contextual:
        el = rpr.makeelement(f"{{{W14}}}cntxtAlts", {f"{{{W14}}}val": "1"})
        rpr.append(el)


def make_docx(path: Path, embed: bool = False) -> None:
    doc = Document()
    doc.add_paragraph("臺灣拼音楷 Word 實測（A＝預設；B＝開啟內容替代字）")
    for title, text in CASES:
        doc.add_paragraph(f"{title}（預期：{EXPECT[title]}）")
        for label, ctx in (("A", False), ("B", True)):
            p = doc.add_paragraph()
            p.add_run(f"{label} ").font.size = Pt(12)
            set_run_font(p.add_run(text), contextual=ctx)
    if embed:
        settings = doc.settings.element
        for tag in ("w:embedTrueTypeFonts",):
            el = OxmlElement(tag)
            settings.insert(0, el)
    doc.save(path)


def make_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(9.2), Inches(6.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (title, text) in enumerate(CASES):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = text
        run.font.size = PPt(26)
        run.font.name = FONT
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.makeelement(qn(tag), {"typeface": FONT})
            rpr.append(el)
    prs.save(path)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 80
    for i, (title, text) in enumerate(CASES, 1):
        ws.cell(i, 1, title)
        c = ws.cell(i, 2, text)
        c.font = XlFont(name=FONT, size=20)
        ws.row_dimensions[i].height = 48
    wb.save(path)


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "out"
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "word-test.docx")
    make_docx(out / "embed.docx", embed=True)
    make_pptx(out / "ppt-test.pptx")
    make_xlsx(out / "excel-test.xlsx")
    print(f"fixtures → {out}")


if __name__ == "__main__":
    main()
